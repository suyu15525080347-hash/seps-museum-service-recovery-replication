# -*- coding: utf-8 -*-
"""Build SEPS manuscript v0.4 (figures integrated) from frozen v0.3 md/docx. No numerical result is changed."""
import sys, os, re, copy, json, hashlib, shutil
import docx
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, r'D:\seps_fig_work')
from ms_edits_v04 import PH, EDITS, NEW_PARAS
from captions_v04 import CAPTIONS, SUPP
S34 = r'G:\我的云端硬盘\艺术计算研究｜Cities+SEPS SSCI投稿｜2026-10\05_阶段审计与城市匹配\34_SEPS_manuscript_prototype_and_internal_desk_reject_v1'
S35 = r'G:\我的云端硬盘\艺术计算研究｜Cities+SEPS SSCI投稿｜2026-10\05_阶段审计与城市匹配\35_SEPS_figures_v2_city_atlas_unified_system'
OUT = r'D:\seps_fig_work\ms'; os.makedirs(OUT, exist_ok=True)
FIG = os.path.join(S35, 'figures')
# ---------------- markdown
md = open(os.path.join(S34, 'SEPS_full_manuscript_v0.3_submission_ready.md'), encoding='utf-8').read()
for old, new in PH.items():
    o = f'**{old}**'; assert md.count(o) == 1, o
    md = md.replace(o, new if old.startswith('[Figure 0') else f'**{new}**')
for old, new in EDITS:
    assert md.count(old) == 1, ('EDIT', old[:60], md.count(old)); md = md.replace(old, new)
for anchor, para in NEW_PARAS:
    assert md.count(anchor) == 1, anchor; md = md.replace(anchor, anchor + '\n\n' + para)
cap_md = '\n\n## Figure captions\n\n' + '\n\n'.join(f'**{t}** {b}' for _, _, t, b in CAPTIONS) + f'\n\n**{SUPP[2]}**\n'
md = md.replace('\n## References', cap_md + '\n## References') if '## Figure captions' not in md else md
open(os.path.join(OUT, 'SEPS_full_manuscript_v0.4_figures_integrated.md'), 'w', encoding='utf-8').write(md)
# ---------------- docx
D = docx.Document(os.path.join(S34, 'SEPS_manuscript_v0.3_submission_ready.docx'))
def set_text(p, text, bold=None):
    runs = p.runs
    if not runs: p.add_run(text); return
    runs[0].text = text
    if bold is not None: runs[0].bold = bold
    for r in runs[1:]: r._element.getparent().remove(r._element)
def insert_after(p, text, bold=None):
    new = copy.deepcopy(p._element); p._element.addnext(new)
    np_ = docx.text.paragraph.Paragraph(new, p._parent); set_text(np_, text, bold); return np_
def uniform(p):
    f = {(bool(r.bold), bool(r.italic), r.font.superscript, r.font.subscript) for r in p.runs if r.text.strip()}
    return len(f) <= 1
log = []
paras = list(D.paragraphs)
for old, new in PH.items():
    hit = [p for p in paras if p.text.strip() == old]; assert len(hit) == 1, old
    set_text(hit[0], new, bold=False if old.startswith('[Figure 0') else True); log.append(('PH', old[:40]))
for old, new in EDITS:
    hit = [p for p in D.paragraphs if old in p.text]; assert len(hit) == 1, ('DOCX', old[:60], len(hit))
    p = hit[0]
    if not uniform(p): log.append(('WARN_nonuniform_runs', old[:50]))
    parts = p.text.replace(old, new).split('\n\n')
    set_text(p, parts[0])
    cur = p
    for part in parts[1:]:
        isb = part.startswith('**'); cur = insert_after(cur, part.strip('*'), bold=True if isb else False)
    log.append(('EDIT', old[:40]))
for anchor, para in NEW_PARAS:
    hit = [p for p in D.paragraphs if anchor in p.text]; assert len(hit) == 1, anchor
    insert_after(hit[0], para, bold=False); log.append(('NEW', anchor[:40]))
# rebuild figure block: delete old Figure 0-4 captions/images (after Table 5), append new
body = D.element.body; ps = list(D.paragraphs)
i0 = [i for i, p in enumerate(ps) if p.text.startswith('Figure 0. Source-audited workflow')][0]
tail = ps[i0:]
for p in tail:
    if p._element.getparent() is body: body.remove(p._element)
sect = body[-1]
TEXT_W = D.sections[0].page_width - D.sections[0].left_margin - D.sections[0].right_margin
cap_style_src = [p for p in D.paragraphs if p.text.startswith('Table 1. Official data sources')][0]
def add_cap(title, text):
    p = D.add_paragraph(); p.style = cap_style_src.style
    r = p.add_run(title + ' '); r.bold = True; r.font.size = Pt(10)
    if text: r2 = p.add_run(text); r2.font.size = Pt(9)
    p.paragraph_format.keep_with_next = True
    return p
from docx.enum.section import WD_SECTION, WD_ORIENT
s0 = D.sections[0]
ns = D.add_section(WD_SECTION.NEW_PAGE)          # figures in a portrait section (tables section stays landscape)
ns.orientation = WD_ORIENT.PORTRAIT; ns.page_width, ns.page_height = s0.page_width, s0.page_height
for a in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(ns, a, getattr(s0, a))
TEXT_W = s0.page_width - s0.left_margin - s0.right_margin
for n, (key, fn, title, text) in enumerate(CAPTIONS):
    pi = D.add_paragraph(); pi.alignment = WD_ALIGN_PARAGRAPH.CENTER; pi.paragraph_format.keep_with_next = True
    if n: pi.paragraph_format.page_break_before = True
    pi.add_run().add_picture(os.path.join(FIG, fn), width=Mm(158) if key == 'Figure 1' else TEXT_W)
    add_cap(title, text)
D.save(os.path.join(OUT, 'SEPS_manuscript_v0.4_figures_integrated.docx'))
# supplementary
S = docx.Document(); S.sections[0].page_width, S.sections[0].page_height = D.sections[0].page_width, D.sections[0].page_height
for a in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(S.sections[0], a, getattr(D.sections[0], a))
S.add_paragraph().add_run('Supplementary material').bold = True
p = S.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(os.path.join(S34, SUPP[1]), width=TEXT_W)
S.add_paragraph().add_run(SUPP[2]).bold = True
S.save(os.path.join(OUT, 'SEPS_supplementary_Fig_S1_workflow_v0.4.docx'))
json.dump(log, open(os.path.join(OUT, 'edit_log.json'), 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
print('built; edits', len(log)); print([l for l in log if l[0].startswith('WARN')])
